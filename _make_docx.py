from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_code_block(doc, code_text):
    for line in code_text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)

def parse_md(md_text, doc):
    lines = md_text.split("\n")
    in_code = False
    code_buf = []
    in_table = False
    table_rows = []

    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue
        if not line.strip():
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            continue
        if line.strip().startswith("|"):
            in_table = True
            if "---" not in line:
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                table_rows.append(cells)
            continue
        else:
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            h = doc.add_heading(line[3:], level=2)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif line.startswith("# "):
            h = doc.add_heading(line[2:], level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif line.strip() == "---":
            doc.add_paragraph("_" * 60)
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.add_run(line)
    if in_table and table_rows:
        _flush_table(doc, table_rows)

def _flush_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j < cols:
                c = table.rows[i].cells[j]
                c.text = cell
                if i == 0:
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.bold = True

with open("LOGIC_FULL.md", "r", encoding="utf-8") as f:
    md = f.read()

doc = Document()
title = doc.add_heading("", 0)
run = title.add_run("AI-трейдер MOEX - Логика работы")
run.font.size = Pt(26)
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sr = sub.add_run("Версия 1.0 - Октябрь 2026")
sr.font.size = Pt(12)
sr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()
parse_md(md, doc)
doc.save("LOGIC.docx")
print("OK - LOGIC.docx создан")
